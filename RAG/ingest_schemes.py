
from pathlib import Path

from openpyxl import load_workbook
from langchain_core.documents import Document
from langchain_chroma import Chroma

from src.embeddings import generate_embedding

RAG_DIR = Path(__file__).resolve().parent
EXCEL_PATH = RAG_DIR / "knowledge_base" / "schemes.xlsx"
CHROMA_PATH = RAG_DIR / "chroma_db"
COLLECTION_NAME = "jansaarthi_schemes"


def load_scheme_documents():
    if not EXCEL_PATH.exists():
        raise FileNotFoundError(f"Excel file not found: {EXCEL_PATH}")

    workbook = load_workbook(
        EXCEL_PATH,
        read_only=True,
        data_only=True
    )

    try:
        if "Scheme Data" not in workbook.sheetnames:
            raise ValueError(
                f"Scheme Data sheet not found. Sheets: {workbook.sheetnames}"
            )

        sheet = workbook["Scheme Data"]
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)

        if not headers:
            raise ValueError("Scheme Data sheet is empty.")

        headers = [
            str(h).strip() if h is not None else ""
            for h in headers
        ]

        required = {"scheme_id", "scheme_name"}
        missing = required - set(headers)

        if missing:
            raise ValueError(f"Missing required columns: {missing}")

        documents = []
        ids = []

        for row in rows:
            record = {
                headers[i]: value
                for i, value in enumerate(row)
                if i < len(headers) and headers[i]
            }

            scheme_id = str(record.get("scheme_id") or "").strip()
            scheme_name = str(record.get("scheme_name") or "").strip()

            if not scheme_id or not scheme_name:
                continue

            # Convert every available Excel field into searchable text.
            content = "\n".join(
                f"{key}: {value}"
                for key, value in record.items()
                if value is not None and str(value).strip()
            )

            documents.append(
                Document(
                    page_content=content,
                    metadata={
                        "scheme_id": scheme_id,
                        "scheme_name": scheme_name,
                        "category": str(record.get("category") or ""),
                        "source_platform": str(
                            record.get("source_platform") or ""
                        ),
                    },
                )
            )

            ids.append(f"scheme-{scheme_id}")

        if not documents:
            raise ValueError("No valid scheme records found.")

        if len(ids) != len(set(ids)):
            raise ValueError("Duplicate scheme IDs found in Excel.")

        return documents, ids

    finally:
        workbook.close()


def main():
    documents, ids = load_scheme_documents()

    vector_store = Chroma(
        collection_name=COLLECTION_NAME,
        embedding_function=generate_embedding(),
        persist_directory=str(CHROMA_PATH),
    )

    # Replace the same IDs on reruns without touching other collections.
    vector_store.delete(ids=ids)
    vector_store.add_documents(documents=documents, ids=ids)

    print(f"Indexed schemes: {len(documents)}")
    print(f"Collection: {COLLECTION_NAME}")
    print(f"Chroma directory: {CHROMA_PATH}")

    results = vector_store.similarity_search(
        "government schemes for small business and entrepreneurs",
        k=5,
    )

    print("\nRetrieval test:")
    for doc in results:
        print("-", doc.metadata.get("scheme_name"))


if __name__ == "__main__":
    main()
